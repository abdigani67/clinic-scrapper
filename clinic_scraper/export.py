"""Export leads to CSV and Excel."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import LEAD_FIELDS, Lead


def to_csv(leads: Iterable[Lead], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=LEAD_FIELDS)
        writer.writeheader()
        for lead in leads:
            writer.writerow(lead.as_row())
    return path


def to_excel(leads: Iterable[Lead], path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F46E5")

    ws.append([f.replace("_", " ").title() for f in LEAD_FIELDS])
    for col, _ in enumerate(LEAD_FIELDS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill

    leads = list(leads)
    for lead in leads:
        row = lead.as_row()
        ws.append([row[f] for f in LEAD_FIELDS])

    # Auto-size columns based on content (capped so they stay readable).
    for col, field_name in enumerate(LEAD_FIELDS, start=1):
        values = [str(lead.as_row()[field_name]) for lead in leads]
        width = max([len(field_name)] + [len(v) for v in values]) + 2
        ws.column_dimensions[get_column_letter(col)].width = min(width, 50)

    ws.freeze_panes = "A2"
    wb.save(path)
    return path
